"""Tests for check_unicode.py."""

import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

from check_unicode import check


def test_clean_doc_no_issues(sample_docx):
    doc = Document(str(sample_docx))
    issues = check(doc)
    assert isinstance(issues, list)
    # Our test doc should be clean
    assert len(issues) == 0


def test_detect_zero_width_space(tmp_path):
    doc = Document()
    doc.add_paragraph("Hello​World")  # Zero-width space
    path = tmp_path / "zwsp.docx"
    doc.save(str(path))

    doc2 = Document(str(path))
    issues = check(doc2)
    assert len(issues) > 0
    assert issues[0]["codepoint"] == "U+200B"
