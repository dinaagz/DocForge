"""Tests for compare_docx.py."""

import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

from compare_docx import compare


def test_compare_identical(sample_docx, tmp_path):
    """Identical documents produce no diffs."""
    doc1 = Document(str(sample_docx))
    doc2 = Document(str(sample_docx))
    report = compare(doc1, doc2)
    assert report["total_diffs"] == 0


def test_compare_style_change(sample_docx, tmp_path):
    """Style-only changes are classified as AUTHORIZED_STRUCTURAL."""
    doc1 = Document(str(sample_docx))
    doc2 = Document(str(sample_docx))

    # Change a heading style
    for p in doc2.paragraphs:
        if p.style and p.style.name == "Heading 1":
            p.style = doc2.styles["Heading 2"]
            break

    report = compare(doc1, doc2)
    structural = report["by_classification"].get("AUTHORIZED_STRUCTURAL", 0)
    assert structural > 0


def test_compare_text_change(sample_docx, tmp_path):
    """Text changes are detected."""
    doc1 = Document(str(sample_docx))
    doc2 = Document(str(sample_docx))

    # Modify some text
    for p in doc2.paragraphs:
        if len(p.text) > 20:
            for run in p.runs:
                if run.text:
                    run.text = run.text.replace("est", "était", 1)
                    break
            break

    report = compare(doc1, doc2)
    assert report["total_diffs"] > 0
