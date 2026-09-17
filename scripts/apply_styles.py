#!/usr/bin/env python3
"""
apply_styles.py — Apply formatting rules from config to a DOCX.

Sets: font, size, margins, alignment, line spacing, heading styles,
pagination, and page numbering. Does NOT modify text content.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, Twips
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

from common import get_logger, load_config, log_event, WORK_DIR

logger = get_logger("apply_styles")

ALIGN_MAP = {
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def apply_formatting(doc: Document, cfg: Dict[str, Any]) -> int:
    """Apply formatting rules. Return number of paragraphs touched."""
    fmt = cfg.get("formatting", {})
    hdg_cfg = cfg.get("headings", {})
    font_family = fmt.get("font", {}).get("family", "Times New Roman")
    font_size = fmt.get("font", {}).get("size", 12)
    alignment = ALIGN_MAP.get(fmt.get("paragraph", {}).get("alignment", "justify"),
                              WD_ALIGN_PARAGRAPH.JUSTIFY)
    line_spacing = fmt.get("paragraph", {}).get("line_spacing", 1.25)

    count = 0
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""

        # Determine if this is a heading
        is_heading = style_name.startswith("Heading")
        heading_num = None
        if is_heading:
            import re
            m = re.search(r"\d+", style_name)
            heading_num = int(m.group()) if m else 1

        # Apply body formatting
        if not is_heading:
            para.alignment = alignment
            pf = para.paragraph_format
            pf.line_spacing = line_spacing

        for run in para.runs:
            run.font.name = font_family
            if is_heading and heading_num:
                key = f"heading_{heading_num}"
                h_cfg = hdg_cfg.get(key, {})
                run.font.size = Pt(h_cfg.get("size", font_size))
                run.font.bold = h_cfg.get("bold", False)
            else:
                run.font.size = Pt(font_size)

        # Page break before for Heading 1
        if is_heading and heading_num == 1:
            h1_cfg = hdg_cfg.get("heading_1", {})
            if h1_cfg.get("page_break_before", True):
                para.paragraph_format.page_break_before = True

        count += 1

    return count


def apply_margins(doc: Document, cfg: Dict[str, Any]) -> None:
    margins = cfg.get("formatting", {}).get("margins", {})
    for section in doc.sections:
        section.top_margin = Cm(margins.get("top_cm", 2.5))
        section.bottom_margin = Cm(margins.get("bottom_cm", 2.5))
        section.left_margin = Cm(margins.get("left_cm", 2.5))
        section.right_margin = Cm(margins.get("right_cm", 2.5))


def add_page_numbers(doc: Document, cfg: Dict[str, Any]) -> None:
    """Add page numbers in footer, centered, excluding first page if configured."""
    pag = cfg.get("pagination", {})
    exclude_first = pag.get("exclude_first_page", True)

    for section in doc.sections:
        if exclude_first:
            section.different_first_page_header_footer = True

        footer = section.footer
        footer.is_linked_to_previous = False
        if footer.paragraphs:
            fp = footer.paragraphs[0]
        else:
            fp = footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Add PAGE field
        run = fp.add_run()
        fldChar1 = OxmlElement("w:fldChar")
        fldChar1.set(qn("w:fldCharType"), "begin")
        run._r.append(fldChar1)

        run2 = fp.add_run()
        instrText = OxmlElement("w:instrText")
        instrText.set(qn("xml:space"), "preserve")
        instrText.text = " PAGE "
        run2._r.append(instrText)

        run3 = fp.add_run()
        fldChar2 = OxmlElement("w:fldChar")
        fldChar2.set(qn("w:fldCharType"), "end")
        run3._r.append(fldChar2)


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply formatting styles to DOCX")
    parser.add_argument("input", type=str, help="Path to DOCX")
    parser.add_argument("--output", type=str, help="Output path (default: overwrite)")
    args = parser.parse_args()

    cfg = load_config()
    inp = Path(args.input)
    if not inp.exists():
        logger.error("File not found: %s", inp)
        return 1

    doc = Document(str(inp))
    count = apply_formatting(doc, cfg)
    apply_margins(doc, cfg)
    add_page_numbers(doc, cfg)

    out = Path(args.output) if args.output else inp
    doc.save(str(out))

    log_event("STYLES_APPLIED", paragraphs=count, output=str(out))
    logger.info("Formatted %d paragraphs -> %s", count, out)
    print(json.dumps({"status": "ok", "paragraphs": count, "output": str(out)}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
