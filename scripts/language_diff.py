#!/usr/bin/env python3
"""
language_diff.py — Apply language corrections to specific paragraphs.

Each correction is logged in corrections.json. Only the exact changes
specified are applied — no full rewrite.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, List

from docx import Document

from common import (
    get_logger, load_state, log_event, now_iso, save_state, text_hash,
)

logger = get_logger("language")


def apply_corrections(doc: Document, corrections: List[Dict[str, Any]]) -> int:
    """Apply targeted corrections. Return number applied."""
    applied = 0
    for corr in corrections:
        pid = corr.get("paragraph_id", "")
        m = re.match(r"P(\d+)", pid)
        if not m:
            logger.warning("Invalid paragraph_id: %s", pid)
            continue
        idx = int(m.group(1))
        if idx >= len(doc.paragraphs):
            logger.warning("Paragraph index %d out of range", idx)
            continue

        para = doc.paragraphs[idx]
        before = corr.get("before", "")
        after = corr.get("after", "")

        if before in para.text:
            # Apply in runs to preserve formatting
            for run in para.runs:
                if before in run.text:
                    run.text = run.text.replace(before, after, 1)
                    applied += 1
                    break
            else:
                # Fallback: concatenate runs, replace, and reset
                full = para.text
                new_text = full.replace(before, after, 1)
                if new_text != full and para.runs:
                    # Put all text in first run, clear others
                    para.runs[0].text = new_text
                    for r in para.runs[1:]:
                        r.text = ""
                    applied += 1
        else:
            logger.warning("Text not found in P%06d: %s", idx, before[:50])

    return applied


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply language corrections to DOCX")
    parser.add_argument("input", type=str, help="Path to DOCX")
    parser.add_argument("--corrections", type=str, required=True, help="Path to corrections JSON")
    parser.add_argument("--output", type=str, help="Output path (default: overwrite)")
    args = parser.parse_args()

    inp = Path(args.input)
    if not inp.exists():
        logger.error("File not found: %s", inp)
        return 1

    with open(args.corrections, "r", encoding="utf-8") as fh:
        corrections = json.load(fh)

    if not isinstance(corrections, list):
        corrections = corrections.get("corrections", [])

    doc = Document(str(inp))
    count = apply_corrections(doc, corrections)

    out = Path(args.output) if args.output else inp
    doc.save(str(out))

    # Update corrections log
    existing = load_state("corrections")
    if not existing:
        existing = {"corrections": [], "total_applied": 0}
    existing["corrections"].extend(corrections)
    existing["total_applied"] = existing.get("total_applied", 0) + count
    existing["last_updated"] = now_iso()
    save_state("corrections", existing)

    log_event("LANGUAGE_CORRECTIONS_APPLIED", count=count, output=str(out))
    logger.info("Applied %d corrections -> %s", count, out)
    print(json.dumps({"status": "ok", "applied": count, "output": str(out)}))
    return 0


if __name__ == "__main__":
    sys.exit(main())
