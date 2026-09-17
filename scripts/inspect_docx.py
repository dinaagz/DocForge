#!/usr/bin/env python3
"""
inspect_docx.py — Inspect a DOCX and produce document_manifest.json

Extracts:
  - paragraph count, estimated pages
  - heading detection (by style + heuristic)
  - tables, images, footnotes/endnotes
  - headers/footers, page breaks
  - existing styles inventory
  - stable paragraph IDs with hashes
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

from common import (
    INPUT_DIR, WORK_DIR, get_logger, input_docx, load_config,
    log_event, now_iso, paragraph_id, save_state, text_hash,
)

logger = get_logger("inspect")

# ── Namespace helpers ──────────────────────────────────────────────
WPNS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _count_page_breaks(doc: Document) -> int:
    """Count explicit page breaks in document XML."""
    count = 0
    for elem in doc.element.iter(f"{WPNS}br"):
        if elem.get(f"{WPNS}type") == "page":
            count += 1
    return count


def _count_images(doc: Document) -> int:
    """Count inline + floating images via relationships."""
    count = 0
    try:
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                count += 1
    except Exception:
        pass
    return count


def _detect_footnotes(doc: Document) -> int:
    count = 0
    for elem in doc.element.iter(f"{WPNS}footnoteReference"):
        count += 1
    return count


def _detect_endnotes(doc: Document) -> int:
    count = 0
    for elem in doc.element.iter(f"{WPNS}endnoteReference"):
        count += 1
    return count


def _has_header_footer(doc: Document) -> Dict[str, bool]:
    has_header = False
    has_footer = False
    for section in doc.sections:
        if section.header and section.header.paragraphs:
            if any(p.text.strip() for p in section.header.paragraphs):
                has_header = True
        if section.footer and section.footer.paragraphs:
            if any(p.text.strip() for p in section.footer.paragraphs):
                has_footer = True
    return {"header": has_header, "footer": has_footer}


def _collect_styles(doc: Document) -> List[str]:
    seen = set()
    for p in doc.paragraphs:
        if p.style and p.style.name:
            seen.add(p.style.name)
    return sorted(seen)


def inspect(docx_path: Path) -> Dict[str, Any]:
    logger.info("Inspecting %s", docx_path)
    doc = Document(str(docx_path))

    paragraphs_info: List[Dict[str, Any]] = []
    headings: List[Dict[str, Any]] = []
    current_section = None

    for idx, para in enumerate(doc.paragraphs):
        pid = paragraph_id(idx)
        txt = para.text
        style_name = para.style.name if para.style else "Normal"

        is_heading = False
        heading_level = None
        if style_name and style_name.startswith("Heading"):
            is_heading = True
            m = re.search(r"\d+", style_name)
            heading_level = int(m.group()) if m else 1

        # Heuristic: short bold paragraphs may be headings
        heuristic_heading = False
        if not is_heading and txt.strip():
            runs = para.runs
            if runs and all(r.bold for r in runs if r.text.strip()):
                if len(txt.strip()) < 120:
                    heuristic_heading = True

        ptype = "heading" if is_heading else ("heuristic_heading" if heuristic_heading else "body")

        p_entry = {
            "id": pid,
            "index": idx,
            "text_preview": txt[:200],
            "text_hash": text_hash(txt),
            "char_count": len(txt),
            "style": style_name,
            "type": ptype,
        }
        if is_heading:
            p_entry["heading_level"] = heading_level
        paragraphs_info.append(p_entry)

        if is_heading or heuristic_heading:
            headings.append({
                "id": pid,
                "index": idx,
                "text": txt.strip(),
                "style": style_name,
                "detected_level": heading_level,
                "heuristic": heuristic_heading,
            })

    # Estimate pages (rough: ~40 lines per page, but we count chars)
    total_chars = sum(p["char_count"] for p in paragraphs_info)
    estimated_pages = max(1, round(total_chars / 2000))

    hf = _has_header_footer(doc)

    manifest = {
        "source_file": docx_path.name,
        "inspected_at": now_iso(),
        "statistics": {
            "paragraph_count": len(doc.paragraphs),
            "estimated_pages": estimated_pages,
            "total_characters": total_chars,
            "table_count": len(doc.tables),
            "image_count": _count_images(doc),
            "footnote_count": _detect_footnotes(doc),
            "endnote_count": _detect_endnotes(doc),
            "page_break_count": _count_page_breaks(doc),
            "section_count": len(doc.sections),
            "heading_count": len(headings),
            "has_header": hf["header"],
            "has_footer": hf["footer"],
        },
        "styles_used": _collect_styles(doc),
        "headings": headings,
        "paragraphs": paragraphs_info,
    }

    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description="Inspect a DOCX file")
    parser.add_argument("--input", type=str, help="Path to DOCX (default: from config)")
    args = parser.parse_args()

    cfg = load_config()
    docx_path = Path(args.input) if args.input else input_docx(cfg)

    if not docx_path.exists():
        logger.error("File not found: %s", docx_path)
        print(json.dumps({"error": f"File not found: {docx_path}"}))
        return 1

    manifest = inspect(docx_path)

    save_state("document_manifest", manifest)
    # Also save a copy under work/inspection/
    out = WORK_DIR / "inspection" / "document_manifest.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8") as fh:
        json.dump(manifest, fh, ensure_ascii=False, indent=2)

    log_event("INSPECTION_COMPLETE",
              paragraphs=manifest["statistics"]["paragraph_count"],
              headings=manifest["statistics"]["heading_count"])

    logger.info("Inspection complete — %d paragraphs, %d headings detected",
                manifest["statistics"]["paragraph_count"],
                manifest["statistics"]["heading_count"])

    print(json.dumps({"status": "ok",
                       "paragraphs": manifest["statistics"]["paragraph_count"],
                       "headings": manifest["statistics"]["heading_count"],
                       "estimated_pages": manifest["statistics"]["estimated_pages"]},
                      indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
