#!/usr/bin/env python3
"""
compare_docx.py — Compare two DOCX files paragraph-by-paragraph.

Produces a diff report classifying each change as:
  AUTHORIZED_STRUCTURAL | AUTHORIZED_LANGUAGE | UNAUTHORIZED | MANUAL_REVIEW
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Dict, List

from docx import Document

from common import (
    get_logger, log_event, now_iso, paragraph_id, text_hash, save_state,
)

logger = get_logger("compare")


def compare(original: Document, modified: Document) -> Dict[str, Any]:
    orig_paras = [(p.text, p.style.name if p.style else "") for p in original.paragraphs]
    mod_paras = [(p.text, p.style.name if p.style else "") for p in modified.paragraphs]

    diffs: List[Dict[str, Any]] = []
    max_len = max(len(orig_paras), len(mod_paras))

    for i in range(max_len):
        pid = paragraph_id(i)

        if i >= len(orig_paras):
            diffs.append({
                "id": pid, "type": "ADDED",
                "classification": "UNAUTHORIZED",
                "modified_text": mod_paras[i][0][:200],
            })
            continue

        if i >= len(mod_paras):
            diffs.append({
                "id": pid, "type": "DELETED",
                "classification": "UNAUTHORIZED",
                "original_text": orig_paras[i][0][:200],
            })
            continue

        orig_text, orig_style = orig_paras[i]
        mod_text, mod_style = mod_paras[i]

        if orig_text == mod_text and orig_style == mod_style:
            continue  # identical

        change: Dict[str, Any] = {"id": pid}

        if orig_text == mod_text and orig_style != mod_style:
            change["type"] = "STYLE_CHANGE"
            change["classification"] = "AUTHORIZED_STRUCTURAL"
            change["original_style"] = orig_style
            change["modified_style"] = mod_style
        elif orig_text != mod_text:
            # Check if it's a minor language fix vs substantive change
            orig_stripped = orig_text.strip().lower()
            mod_stripped = mod_text.strip().lower()

            # Simple heuristic: if length differs by >20% it's suspicious
            len_ratio = abs(len(orig_text) - len(mod_text)) / max(len(orig_text), 1)

            if len_ratio > 0.2:
                change["classification"] = "MANUAL_REVIEW"
            else:
                change["classification"] = "AUTHORIZED_LANGUAGE"

            change["type"] = "TEXT_MODIFIED"
            change["original_preview"] = orig_text[:200]
            change["modified_preview"] = mod_text[:200]
            change["original_hash"] = text_hash(orig_text)
            change["modified_hash"] = text_hash(mod_text)
            change["length_change_ratio"] = round(len_ratio, 3)

        diffs.append(change)

    report = {
        "compared_at": now_iso(),
        "original_paragraphs": len(orig_paras),
        "modified_paragraphs": len(mod_paras),
        "total_diffs": len(diffs),
        "by_classification": {},
        "diffs": diffs,
    }

    # Count by classification
    for d in diffs:
        cls = d.get("classification", "UNKNOWN")
        report["by_classification"][cls] = report["by_classification"].get(cls, 0) + 1

    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="Compare two DOCX files")
    parser.add_argument("original", help="Path to original DOCX")
    parser.add_argument("modified", help="Path to modified DOCX")
    parser.add_argument("--output", help="Output JSON path")
    args = parser.parse_args()

    orig_path = Path(args.original)
    mod_path = Path(args.modified)
    for p in (orig_path, mod_path):
        if not p.exists():
            logger.error("File not found: %s", p)
            return 1

    orig_doc = Document(str(orig_path))
    mod_doc = Document(str(mod_path))
    report = compare(orig_doc, mod_doc)

    if args.output:
        Path(args.output).parent.mkdir(parents=True, exist_ok=True)
        with open(args.output, "w", encoding="utf-8") as fh:
            json.dump(report, fh, ensure_ascii=False, indent=2)

    save_state("validation_issues", report)
    log_event("COMPARISON_COMPLETE", diffs=report["total_diffs"],
              by_class=report["by_classification"])

    print(json.dumps({
        "status": "ok",
        "diffs": report["total_diffs"],
        "by_classification": report["by_classification"],
    }, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
